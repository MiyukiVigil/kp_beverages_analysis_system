import pandas as pd
import io
import os
import re
import tempfile

from fpdf import FPDF
from datetime import datetime
from config import TIN_DRINKS
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

CURRENCY_FORMAT = '"RM "#,##0.00'
DATE_FORMAT = "yyyy-mm-dd"
INTEGER_FORMAT = "#,##0"
KG_FORMAT = "0.000"
PERCENT_FORMAT = "0.0%"
CHART_HEIGHT = 13
CHART_WIDTH = 22.5
WIDE_DATE_CHART_WIDTH = 32
BAR_POINT_COLORS = ["4F81BD", "C0504D", "9BBB59", "8064A2", "4BACC6", "F79646", "92A9CF", "D99694", "A9C47F", "B1A0C7"]
BAR_SERIES_COLORS = ["4F81BD", "C0504D", "9BBB59", "8064A2", "4BACC6"]

# ==========================================================
# EXCEL EXPORT (SAFE)
# ==========================================================

def generate_excel(sheets_dict):
    output = io.BytesIO()

    if not sheets_dict:
        sheets_dict = {
            "System Status": pd.DataFrame(
                {"Status": ["No data available to export"]}
            )
        }

    with pd.ExcelWriter(output, engine="openpyxl") as writer:

        for sheet_name, df in sheets_dict.items():

            safe_sheet_name = str(sheet_name).replace("/", "_").replace("\\", "_")[:31]

            if df is None or df.empty:
                df = pd.DataFrame({"Status": ["No data available"]})

            df.to_excel(writer, index=False, sheet_name=safe_sheet_name)

            try:
                worksheet = writer.sheets[safe_sheet_name]

                for idx, col in enumerate(df.columns):
                    series = df[col]

                    max_len = max(
                        (
                            series.astype(str).map(len).max()
                            if not series.empty else 0
                        ),
                        len(str(col))
                    ) + 2

                    col_letter = get_column_letter(idx + 1)

                    worksheet.column_dimensions[col_letter].width = min(max_len, 50)

            except Exception:
                pass

    return output.getvalue()

def generate_sales_analysis_excel(df, menu_dict, settings_dict, date_range_str):
    """
    Builds a multi-sheet analysis workbook for sales review and reconciliation.
    The PDF export stays concise; this workbook carries the audit/detail layer.
    """
    df = _prepare_sales_analysis_df(df)
    if df.empty:
        return generate_excel({
            "Dashboard": pd.DataFrame({"Status": ["No sales data available for this period"]})
        })

    daily_summary = _build_daily_summary(df)
    product_performance = _build_product_performance(df)
    payment_breakdown = _build_payment_breakdown(df)
    coffee_reconciliation = _build_coffee_reconciliation(df)
    modifier_analysis = _build_modifier_analysis(df)
    menu_reference = _build_menu_reference(menu_dict, settings_dict)
    checks = _build_checks(df, daily_summary)

    raw_columns = [
        "Date", "Drink Base", "Drink Type", "Drink Profile", "Qty",
        "Base Price", "Unit Price", "Takeaway Fee",
        "Is Takeaway", "Is Kosong", "Payment Method",
        "Revenue", "Cash Revenue", "QR Revenue", "Expected Coffee (kg)"
    ]
    raw_ledger = df[[col for col in raw_columns if col in df.columns]].copy()

    sheets = {
        "Dashboard": _build_dashboard(
            df,
            daily_summary,
            product_performance,
            coffee_reconciliation,
            date_range_str
        ),
        "Daily Summary": daily_summary,
        "Product Performance": product_performance,
        "Payment Breakdown": payment_breakdown,
        "Coffee Reconciliation": coffee_reconciliation,
        "Modifier Analysis": modifier_analysis,
        "Raw Ledger": raw_ledger,
        "Menu Reference": menu_reference,
        "Checks": checks,
    }

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, sheet_df in sheets.items():
            if sheet_df is None or sheet_df.empty:
                sheet_df = pd.DataFrame({"Status": ["No data available for this section"]})
            sheet_df.to_excel(writer, index=False, sheet_name=sheet_name)

        workbook = writer.book
        for sheet_name in sheets:
            worksheet = workbook[sheet_name]
            _format_sheet(worksheet)
            _format_sheet_as_table(worksheet)
            _autosize_columns(worksheet)

        _format_dashboard(workbook["Dashboard"])
        _add_workbook_charts(workbook)

    return output.getvalue()

def _prepare_sales_analysis_df(df):
    if df is None or df.empty:
        return pd.DataFrame()

    df = df.copy()
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df = df.dropna(subset=["Date"])
    if df.empty:
        return df

    for col in ["Drink Base", "Drink Type", "Drink Profile", "Payment Method"]:
        if col not in df.columns:
            df[col] = ""
        df[col] = df[col].fillna("").astype(str)

    for col in [
        "Qty", "Base Price", "Unit Price", "Takeaway Fee", "Revenue",
        "Cash Revenue", "QR Revenue", "Takeaway", "Kosong",
        "Expected Coffee (kg)", "Day Actual Coffee (kg)", "Day Expected Coffee (kg)"
    ]:
        if col not in df.columns:
            df[col] = 0.0
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0.0)

    if "Cash Revenue" not in df.columns:
        df["Cash Revenue"] = df["Revenue"] - df["QR Revenue"]
    else:
        df["Cash Revenue"] = df["Revenue"] - df["QR Revenue"]

    blank_payment = df["Payment Method"] == ""
    df.loc[blank_payment, "Payment Method"] = df.loc[blank_payment, "QR Revenue"].apply(
        lambda value: "QR" if value > 0 else "Cash"
    )

    if "Is Takeaway" not in df.columns:
        df["Is Takeaway"] = df["Takeaway"] > 0
    else:
        df["Is Takeaway"] = df["Is Takeaway"].fillna(False).astype(bool)

    if "Is Kosong" not in df.columns:
        df["Is Kosong"] = df["Kosong"] > 0
    else:
        df["Is Kosong"] = df["Is Kosong"].fillna(False).astype(bool)

    missing_unit_price = df["Unit Price"] == 0
    qty_denominator = df["Qty"].where(df["Qty"] != 0)
    df.loc[missing_unit_price, "Unit Price"] = (
        df.loc[missing_unit_price, "Revenue"] / qty_denominator.loc[missing_unit_price]
    ).fillna(0)

    missing_base_price = df["Base Price"] == 0
    df.loc[missing_base_price, "Base Price"] = df.loc[missing_base_price, "Unit Price"]
    return df

def _build_dashboard(df, daily_summary, product_performance, coffee_reconciliation, date_range_str):
    gross_revenue = df["Revenue"].sum()
    cash_revenue = df["Cash Revenue"].sum()
    qr_revenue = df["QR Revenue"].sum()
    total_cups = df["Qty"].sum()
    top_qty = product_performance.sort_values("Qty Sold", ascending=False).iloc[0]
    top_revenue = product_performance.sort_values("Gross Revenue", ascending=False).iloc[0]
    expected_coffee = coffee_reconciliation["Expected Coffee kg"].sum()
    actual_coffee = coffee_reconciliation["Actual Coffee kg"].sum()
    variance = actual_coffee - expected_coffee

    rows = [
        ["Report Period", date_range_str, "", "Selected export range"],
        ["Generated On", datetime.now().strftime("%Y-%m-%d %H:%M"), "", ""],
        ["Days Included", daily_summary["Date"].nunique(), "days", ""],
        ["Transaction Rows", len(df), "rows", "Ledger rows, not unique receipts"],
        ["Total Cups Sold", total_cups, "cups", ""],
        ["Gross Revenue", gross_revenue, "RM", ""],
        ["Cash Revenue", cash_revenue, "RM", ""],
        ["QR Revenue", qr_revenue, "RM", ""],
        ["Cashless (QR) Percentage", _safe_scalar_divide(qr_revenue, gross_revenue), "%", ""],
        ["Average Revenue Per Cup", _safe_scalar_divide(gross_revenue, total_cups), "RM/cup", ""],
        ["Top Item by Quantity", top_qty["Drink Profile"], "", f"{top_qty['Qty Sold']:.0f} cups"],
        ["Top Item by Revenue", top_revenue["Drink Profile"], "", f"RM {top_revenue['Gross Revenue']:.2f}"],
        ["Expected Coffee Used", expected_coffee, "kg", ""],
        ["Actual Coffee Used", actual_coffee, "kg", ""],
        ["Coffee Variance", variance, "kg", _coffee_status(expected_coffee, actual_coffee, variance)],
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value", "Unit", "Notes"])

def _build_daily_summary(df):
    work = df.copy()
    work["Report Date"] = work["Date"].dt.date
    daily = work.groupby("Report Date").agg(**{
        "Transaction Rows": ("Qty", "count"),
        "Total Cups": ("Qty", "sum"),
        "Gross Revenue": ("Revenue", "sum"),
        "Cash Revenue": ("Cash Revenue", "sum"),
        "QR Revenue": ("QR Revenue", "sum"),
        "Takeaway Cups": ("Takeaway", "sum"),
        "Kosong Cups": ("Kosong", "sum"),
        "Expected Coffee kg": ("Day Expected Coffee (kg)", "max"),
        "Actual Coffee kg": ("Day Actual Coffee (kg)", "max"),
    }).reset_index().rename(columns={"Report Date": "Date"})

    daily["Cashless (QR) Percentage"] = _safe_series_divide(daily["QR Revenue"], daily["Gross Revenue"])
    daily["Revenue Per Cup"] = _safe_series_divide(daily["Gross Revenue"], daily["Total Cups"])
    daily["Coffee Variance kg"] = daily["Actual Coffee kg"] - daily["Expected Coffee kg"]
    daily["Variance Status"] = daily.apply(
        lambda row: _coffee_status(row["Expected Coffee kg"], row["Actual Coffee kg"], row["Coffee Variance kg"]),
        axis=1
    )

    return daily[[
        "Date", "Transaction Rows", "Total Cups", "Gross Revenue", "Cash Revenue",
        "QR Revenue", "Cashless (QR) Percentage", "Takeaway Cups", "Kosong Cups",
        "Expected Coffee kg", "Actual Coffee kg", "Coffee Variance kg",
        "Revenue Per Cup", "Variance Status"
    ]]

def _build_product_performance(df):
    products = df.groupby(["Drink Profile", "Drink Base", "Drink Type"]).agg(**{
        "Qty Sold": ("Qty", "sum"),
        "Gross Revenue": ("Revenue", "sum"),
        "Cash Revenue": ("Cash Revenue", "sum"),
        "QR Revenue": ("QR Revenue", "sum"),
        "Expected Coffee kg": ("Expected Coffee (kg)", "sum"),
    }).reset_index()

    products["Qty Share"] = _safe_series_divide(products["Qty Sold"], products["Qty Sold"].sum())
    products["Revenue Share"] = _safe_series_divide(products["Gross Revenue"], products["Gross Revenue"].sum())
    products["Avg Revenue Per Cup"] = _safe_series_divide(products["Gross Revenue"], products["Qty Sold"])
    return products[[
        "Drink Profile", "Drink Base", "Drink Type", "Qty Sold",
        "Gross Revenue", "Cash Revenue", "QR Revenue", "Qty Share",
        "Revenue Share", "Avg Revenue Per Cup", "Expected Coffee kg"
    ]].sort_values("Gross Revenue", ascending=False)

def _build_payment_breakdown(df):
    work = df.copy()
    work["Report Date"] = work["Date"].dt.date
    work["Cash Rows"] = (work["Payment Method"] == "Cash").astype(int)
    work["QR Rows"] = (work["Payment Method"] == "QR").astype(int)
    work["Cash Cups"] = work["Qty"].where(work["Payment Method"] == "Cash", 0)
    work["QR Cups"] = work["Qty"].where(work["Payment Method"] == "QR", 0)

    payment = work.groupby("Report Date").agg(**{
        "Cash Revenue": ("Cash Revenue", "sum"),
        "QR Revenue": ("QR Revenue", "sum"),
        "Gross Revenue": ("Revenue", "sum"),
        "Cash Rows": ("Cash Rows", "sum"),
        "QR Rows": ("QR Rows", "sum"),
        "Cash Cups": ("Cash Cups", "sum"),
        "QR Cups": ("QR Cups", "sum"),
    }).reset_index().rename(columns={"Report Date": "Date"})

    payment["Cash Share"] = _safe_series_divide(payment["Cash Revenue"], payment["Gross Revenue"])
    payment["Cashless (QR) Percentage"] = _safe_series_divide(payment["QR Revenue"], payment["Gross Revenue"])
    return payment[[
        "Date", "Cash Revenue", "QR Revenue", "Gross Revenue", "Cash Share",
        "Cashless (QR) Percentage", "Cash Rows", "QR Rows", "Cash Cups", "QR Cups"
    ]]

def _build_coffee_reconciliation(df):
    work = df.copy()
    work["Report Date"] = work["Date"].dt.date
    drink_base = work["Drink Base"].astype(str)
    drink_type = work["Drink Type"].astype(str)
    is_kopi = drink_base.str.contains("Kopi", na=False)

    work["Hot Kopi Cups"] = work["Qty"].where(is_kopi & drink_type.eq("Hot"), 0)
    work["Cold Kopi Cups"] = work["Qty"].where(is_kopi & drink_type.eq("Cold"), 0)

    coffee = work.groupby("Report Date").agg(**{
        "Hot Kopi Cups": ("Hot Kopi Cups", "sum"),
        "Cold Kopi Cups": ("Cold Kopi Cups", "sum"),
        "Expected Coffee kg": ("Day Expected Coffee (kg)", "max"),
        "Actual Coffee kg": ("Day Actual Coffee (kg)", "max"),
    }).reset_index().rename(columns={"Report Date": "Date"})

    coffee["Variance kg"] = coffee["Actual Coffee kg"] - coffee["Expected Coffee kg"]
    coffee["Variance %"] = _safe_series_divide(coffee["Variance kg"], coffee["Expected Coffee kg"])
    coffee["Status"] = coffee.apply(
        lambda row: _coffee_status(row["Expected Coffee kg"], row["Actual Coffee kg"], row["Variance kg"]),
        axis=1
    )
    coffee["Notes"] = ""
    return coffee

def _build_modifier_analysis(df):
    work = df.copy()
    work["Report Date"] = work["Date"].dt.date
    work["Hot Cups"] = work["Qty"].where(work["Drink Type"].eq("Hot"), 0)
    work["Cold Cups"] = work["Qty"].where(work["Drink Type"].eq("Cold"), 0)
    work["QR Cups"] = work["Qty"].where(work["Payment Method"].eq("QR"), 0)
    work["Cash Cups"] = work["Qty"].where(work["Payment Method"].eq("Cash"), 0)

    modifiers = work.groupby("Report Date").agg(**{
        "Total Cups": ("Qty", "sum"),
        "Takeaway Cups": ("Takeaway", "sum"),
        "Takeaway Fees": ("Takeaway Fee", "sum"),
        "Kosong Cups": ("Kosong", "sum"),
        "Hot Cups": ("Hot Cups", "sum"),
        "Cold Cups": ("Cold Cups", "sum"),
        "QR Cups": ("QR Cups", "sum"),
        "Cash Cups": ("Cash Cups", "sum"),
    }).reset_index().rename(columns={"Report Date": "Date"})

    modifiers["Takeaway Share"] = _safe_series_divide(modifiers["Takeaway Cups"], modifiers["Total Cups"])
    modifiers["Kosong Share"] = _safe_series_divide(modifiers["Kosong Cups"], modifiers["Total Cups"])
    return modifiers[[
        "Date", "Total Cups", "Takeaway Cups", "Takeaway Share",
        "Takeaway Fees", "Kosong Cups", "Kosong Share", "Hot Cups",
        "Cold Cups", "QR Cups", "Cash Cups"
    ]]

def _build_menu_reference(menu_dict, settings_dict):
    rows = []
    small_g = settings_dict.get("small_cup_g", 10.0)
    big_g = settings_dict.get("big_cup_g", 20.0)

    for drink, drink_data in menu_dict.items():
        for temp in drink_data.get("temperature", []):
            drink_type = temp.get("type", "")
            is_coffee = "Kopi" in drink
            coffee_grams = 0.0
            if is_coffee:
                coffee_grams = small_g if drink_type == "Hot" else big_g
            rows.append({
                "Drink Base": drink,
                "Drink Type": drink_type,
                "Base Price": temp.get("price", 0.0),
                "Takeaway Eligible": drink not in TIN_DRINKS,
                "Is Coffee Item": is_coffee,
                "Coffee g Per Cup": coffee_grams,
                "Coffee kg Per Cup": coffee_grams / 1000.0,
            })

    return pd.DataFrame(rows).sort_values(["Drink Base", "Drink Type"])

def _build_checks(df, daily_summary):
    raw_revenue = df["Revenue"].sum()
    daily_revenue = daily_summary["Gross Revenue"].sum()
    raw_cups = df["Qty"].sum()
    daily_cups = daily_summary["Total Cups"].sum()
    missing_actual = daily_summary[
        (daily_summary["Expected Coffee kg"] > 0) & (daily_summary["Actual Coffee kg"] <= 0)
    ]
    variance_alerts = daily_summary[daily_summary["Variance Status"].isin(["Audit Required", "Review Required"])]

    checks = [
        [
            "Revenue ties from raw ledger to daily summary",
            "PASS" if abs(raw_revenue - daily_revenue) < 0.01 else "REVIEW",
            f"Raw RM {raw_revenue:.2f}; Daily RM {daily_revenue:.2f}"
        ],
        [
            "Cup volume ties from raw ledger to daily summary",
            "PASS" if abs(raw_cups - daily_cups) < 0.01 else "REVIEW",
            f"Raw {raw_cups:.0f}; Daily {daily_cups:.0f}"
        ],
        [
            "Rows with missing drink or type",
            "PASS" if df["Drink Base"].eq("").sum() + df["Drink Type"].eq("").sum() == 0 else "REVIEW",
            f"{df['Drink Base'].eq('').sum() + df['Drink Type'].eq('').sum()} missing values"
        ],
        [
            "Rows with non-positive quantity",
            "PASS" if (df["Qty"] <= 0).sum() == 0 else "REVIEW",
            f"{(df['Qty'] <= 0).sum()} rows"
        ],
        [
            "Rows with zero base price",
            "PASS" if (df["Base Price"] <= 0).sum() == 0 else "REVIEW",
            f"{(df['Base Price'] <= 0).sum()} rows"
        ],
        [
            "Coffee days missing actual usage",
            "PASS" if missing_actual.empty else "REVIEW",
            f"{len(missing_actual)} days"
        ],
        [
            "Coffee variance outside tolerance",
            "PASS" if variance_alerts.empty else "REVIEW",
            f"{len(variance_alerts)} days above 0.100 kg tolerance"
        ],
    ]
    return pd.DataFrame(checks, columns=["Check", "Result", "Detail"])

def _safe_series_divide(numerator, denominator):
    if not isinstance(denominator, pd.Series):
        denominator = pd.Series([denominator] * len(numerator), index=numerator.index)
    result = numerator.divide(denominator.where(denominator != 0))
    return result.fillna(0)

def _safe_scalar_divide(numerator, denominator):
    if denominator == 0 or pd.isna(denominator):
        return 0
    return numerator / denominator

def _coffee_status(expected, actual, variance):
    if pd.isna(actual) or actual <= 0:
        return "Missing Actual" if expected > 0 else "No Coffee Data"
    if expected <= 0:
        return "Audit Required" if actual > 0 else "No Coffee Data"
    if variance > 0.1:
        return "Audit Required"
    if variance < -0.1:
        return "Review Required"
    return "OK"

def _format_sheet(worksheet):
    worksheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            header = worksheet.cell(row=1, column=cell.column).value or ""
            _apply_number_format(cell, str(header))
            cell.alignment = Alignment(vertical="top")

def _apply_number_format(cell, header):
    lower_header = header.lower()
    if lower_header == "date":
        cell.number_format = DATE_FORMAT
    elif "kg" in lower_header:
        cell.number_format = KG_FORMAT
    elif "share" in lower_header or "percentage" in lower_header or "%" in lower_header:
        cell.number_format = PERCENT_FORMAT
    elif (
        "revenue" in lower_header
        or "price" in lower_header
        or lower_header.endswith("fee")
        or lower_header.endswith("fees")
    ):
        cell.number_format = CURRENCY_FORMAT
    elif "cups" in lower_header or "rows" in lower_header or "qty" in lower_header or "count" in lower_header:
        cell.number_format = INTEGER_FORMAT

def _format_dashboard(worksheet):
    worksheet.freeze_panes = "A2"
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 24
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 42

    for row in worksheet.iter_rows(min_row=2):
        metric = row[0].value
        value_cell = row[1]
        if metric in {"Gross Revenue", "Cash Revenue", "QR Revenue", "Average Revenue Per Cup"}:
            value_cell.number_format = CURRENCY_FORMAT
        elif metric == "Cashless (QR) Percentage":
            value_cell.number_format = PERCENT_FORMAT
        elif metric in {"Expected Coffee Used", "Actual Coffee Used", "Coffee Variance"}:
            value_cell.number_format = KG_FORMAT
        elif metric in {"Days Included", "Transaction Rows", "Total Cups Sold"}:
            value_cell.number_format = INTEGER_FORMAT

def _format_sheet_as_table(worksheet):
    if worksheet.max_row < 1 or worksheet.max_column < 1:
        return

    ref = f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    if worksheet.max_row < 2:
        worksheet.auto_filter.ref = ref
        return

    table = Table(displayName=_excel_table_name(worksheet.title), ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)

def _excel_table_name(sheet_name):
    cleaned = re.sub(r"[^A-Za-z0-9_]", "", sheet_name)
    if not cleaned:
        cleaned = "Report"
    if cleaned[0].isdigit():
        cleaned = f"T{cleaned}"
    return f"{cleaned[:240]}Table"

def _autosize_columns(worksheet):
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(max_len + 2, 10), 45)

def _add_workbook_charts(workbook):
    _add_line_chart(
        workbook["Daily Summary"],
        title="Daily gross revenue trend (RM)",
        category_header="Date",
        value_headers=["Gross Revenue"],
        y_axis_title="Revenue (RM)",
        x_axis_title="Date",
        position=_chart_position_below_table(workbook["Daily Summary"]),
        y_axis_number_format=CURRENCY_FORMAT,
    )
    _add_bar_chart(
        workbook["Product Performance"],
        title="Top 10 products by revenue (RM)",
        category_header="Drink Profile",
        value_headers=["Gross Revenue"],
        y_axis_title="Revenue (RM)",
        x_axis_title="Drink",
        position=_chart_position_below_table(workbook["Product Performance"]),
        max_rows=10,
        horizontal=True,
        value_axis_number_format=CURRENCY_FORMAT,
    )
    _add_line_chart(
        workbook["Payment Breakdown"],
        title="Cashless (QR) percentage by date",
        category_header="Date",
        value_headers=["Cashless (QR) Percentage"],
        y_axis_title="Cashless share",
        x_axis_title="Date",
        position=_chart_position_below_table(workbook["Payment Breakdown"]),
        y_axis_number_format="0%",
        y_axis_min=0,
        y_axis_max=1,
        width=WIDE_DATE_CHART_WIDTH,
    )
    _add_bar_chart(
        workbook["Coffee Reconciliation"],
        title="Coffee variance by date (kg)",
        category_header="Date",
        value_headers=["Variance kg"],
        y_axis_title="Variance (kg)",
        x_axis_title="Date",
        position=_chart_position_below_table(workbook["Coffee Reconciliation"]),
        value_axis_number_format=KG_FORMAT,
        blank_zero_values=True,
        x_axis_at_bottom=True,
        layout_y=0,
        layout_height=0.76,
        width=WIDE_DATE_CHART_WIDTH,
    )
    _add_bar_chart(
        workbook["Modifier Analysis"],
        title="Takeaway and kosong cups by date",
        category_header="Date",
        value_headers=["Takeaway Cups", "Kosong Cups"],
        y_axis_title="Cups",
        x_axis_title="Date",
        position=_chart_position_below_table(workbook["Modifier Analysis"]),
        value_axis_number_format=INTEGER_FORMAT,
        blank_zero_values=True,
        layout_y=0,
        layout_height=0.76,
        width=WIDE_DATE_CHART_WIDTH,
    )

def _chart_position_below_table(worksheet):
    return f"A{worksheet.max_row + 3}"

def _add_line_chart(
    worksheet,
    title,
    category_header,
    value_headers,
    position,
    y_axis_title,
    x_axis_title,
    y_axis_number_format=None,
    y_axis_min=None,
    y_axis_max=None,
    width=CHART_WIDTH,
):
    if worksheet.max_row < 3:
        return

    category_col = _column_index(worksheet, category_header)
    value_cols = [_column_index(worksheet, header) for header in value_headers]
    if not category_col or any(col is None for col in value_cols):
        return

    chart = LineChart()
    chart.title = title
    chart.style = 13
    if y_axis_title:
        chart.y_axis.title = y_axis_title
    if x_axis_title:
        chart.x_axis.title = x_axis_title
    chart.legend = None
    _configure_chart_axes(chart)
    _configure_chart_layout(chart)
    chart.x_axis.numFmt = DATE_FORMAT

    if y_axis_number_format:
        chart.y_axis.numFmt = y_axis_number_format
    if y_axis_min is not None:
        chart.y_axis.scaling.min = y_axis_min
    if y_axis_max is not None:
        chart.y_axis.scaling.max = y_axis_max
    for value_col in value_cols:
        data = Reference(worksheet, min_col=value_col, min_row=1, max_row=worksheet.max_row)
        chart.add_data(data, titles_from_data=True)
    categories = Reference(worksheet, min_col=category_col, min_row=2, max_row=worksheet.max_row)
    chart.set_categories(categories)
    chart.height = CHART_HEIGHT
    chart.width = width
    worksheet.add_chart(chart, position)

def _add_bar_chart(
    worksheet,
    title,
    category_header,
    value_headers,
    position,
    y_axis_title,
    x_axis_title,
    max_rows=None,
    horizontal=False,
    value_axis_number_format=None,
    blank_zero_values=False,
    x_axis_at_bottom=False,
    layout_y=0.04,
    layout_height=0.68,
    width=CHART_WIDTH,
):
    if worksheet.max_row < 3:
        return

    max_row = worksheet.max_row if max_rows is None else min(worksheet.max_row, max_rows + 1)
    category_col = _column_index(worksheet, category_header)
    value_cols = [_column_index(worksheet, header) for header in value_headers]
    if not category_col or any(col is None for col in value_cols):
        return

    chart = BarChart()
    _configure_chart_axes(chart)
    _configure_chart_layout(chart, y=layout_y, h=layout_height)

    if horizontal:
        chart.type = "bar"
    else:
        chart.x_axis.numFmt = DATE_FORMAT
    if x_axis_title:
        chart.x_axis.title = x_axis_title
    if y_axis_title:
        chart.y_axis.title = y_axis_title
    if value_axis_number_format:
        chart.y_axis.numFmt = value_axis_number_format
    if x_axis_at_bottom:
        chart.x_axis.crosses = "min"
    chart.style = 10
    chart.title = title
    if len(value_cols) == 1:
        chart.legend = None
    else:
        chart.legend.position = "r"
        chart.legend.overlay = False

    if blank_zero_values:
        category_col, value_cols = _build_hidden_chart_source(worksheet, category_col, value_cols, max_row)
        chart.visible_cells_only = False

    for value_col in value_cols:
        data = Reference(worksheet, min_col=value_col, min_row=1, max_row=max_row)
        chart.add_data(data, titles_from_data=True)
    categories = Reference(worksheet, min_col=category_col, min_row=2, max_row=max_row)
    chart.set_categories(categories)
    _configure_bar_series(chart, worksheet, value_cols, max_row, value_axis_number_format)
    chart.height = CHART_HEIGHT
    chart.width = width
    worksheet.add_chart(chart, position)

def _configure_chart_axes(chart):
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.tickLblPos = "nextTo"
        axis.majorTickMark = "out"
        axis.minorTickMark = "none"

def _configure_chart_layout(chart, y=0.04, h=0.68):
    chart.layout = Layout(
        manualLayout=ManualLayout(
            layoutTarget="inner",
            xMode="factor",
            yMode="factor",
            wMode="factor",
            hMode="factor",
            x=0.04,
            y=y,
            w=0.91,
            h=h,
        )
    )

def _build_hidden_chart_source(worksheet, category_col, value_cols, max_row):
    helper_start_col = worksheet.max_column + 2
    helper_category_col = helper_start_col
    helper_value_cols = [helper_start_col + index + 1 for index in range(len(value_cols))]

    worksheet.cell(row=1, column=helper_category_col).value = worksheet.cell(row=1, column=category_col).value
    for helper_col, value_col in zip(helper_value_cols, value_cols):
        worksheet.cell(row=1, column=helper_col).value = worksheet.cell(row=1, column=value_col).value

    for row_index in range(2, max_row + 1):
        worksheet.cell(row=row_index, column=helper_category_col).value = worksheet.cell(row=row_index, column=category_col).value
        for helper_col, value_col in zip(helper_value_cols, value_cols):
            value = worksheet.cell(row=row_index, column=value_col).value
            worksheet.cell(row=row_index, column=helper_col).value = None if value == 0 else value

    for hidden_col in [helper_category_col, *helper_value_cols]:
        worksheet.column_dimensions[get_column_letter(hidden_col)].hidden = True

    return helper_category_col, helper_value_cols

def _configure_bar_series(chart, worksheet, value_cols, max_row, number_format=None):
    for series_index, series in enumerate(chart.series):
        values = [
            worksheet.cell(row=row_index, column=value_cols[series_index]).value
            for row_index in range(2, max_row + 1)
        ]
        if series_index == 0:
            _configure_bar_data_labels(chart, number_format)
        if len(chart.series) == 1:
            _apply_point_colors(series, values)
        else:
            _apply_series_color(series, BAR_SERIES_COLORS[series_index % len(BAR_SERIES_COLORS)])

def _configure_bar_data_labels(chart, number_format=None):
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.dataLabels.showCatName = False
    chart.dataLabels.showSerName = False
    chart.dataLabels.showLegendKey = False
    chart.dataLabels.showPercent = False
    chart.dataLabels.dLblPos = "outEnd"
    if number_format:
        chart.dataLabels.numFmt = _zero_suppressed_label_format(number_format)

def _apply_point_colors(series, values):
    for point_index, value in enumerate(values):
        if value in (None, 0):
            continue
        point = DataPoint(idx=point_index)
        point.graphicalProperties.solidFill = BAR_POINT_COLORS[point_index % len(BAR_POINT_COLORS)]
        series.dPt.append(point)

def _apply_series_color(series, color):
    series.graphicalProperties.solidFill = color

def _zero_suppressed_label_format(number_format):
    if number_format == CURRENCY_FORMAT:
        return '"RM "#,##0.00;-"RM "#,##0.00;;'
    if number_format == KG_FORMAT:
        return "0.000;-0.000;;"
    if number_format == INTEGER_FORMAT:
        return "#,##0;-#,##0;;"
    return number_format

def _column_index(worksheet, header):
    for cell in worksheet[1]:
        if cell.value == header:
            return cell.column
    return None

class BusinessPDFReport(FPDF):

    def header(self):
        self.set_font("Helvetica", "B", 14)
        self.set_text_color(30, 30, 30)

        self.cell(
            0,
            10,
            "KOPITIAM BUSINESS PERFORMANCE REPORT",
            ln=True,
            align="C"
        )

        self.set_draw_color(200, 200, 200)
        self.line(10, 22, 200, 22)

        self.ln(10)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(120, 120, 120)

        self.cell(
            0,
            10,
            f"Page {self.page_no()}",
            align="C"
        )


# ==========================================================
# KPI BOX
# ==========================================================

def draw_kpi(pdf, label, value):
    pdf.set_font("Helvetica", "B", 10)
    pdf.set_fill_color(245, 245, 245)

    pdf.cell(95, 10, label, border=1, fill=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(95, 10, str(value), border=1, ln=True)


# ==========================================================
# SAFE CHART GENERATOR (FIXED DATE OVERLAP)
# ==========================================================

def create_chart(series, title, chart_type="line"):
    try:
        import matplotlib.pyplot as plt
        import matplotlib.dates as mdates
    except ModuleNotFoundError as exc:
        raise RuntimeError("PDF chart generation requires matplotlib to be installed.") from exc

    temp = tempfile.NamedTemporaryFile(delete=False, suffix=".png")

    plt.figure(figsize=(10, 4))

    if series is None or len(series) == 0 or series.sum() == 0:
        plt.text(0.5, 0.5, "No Data Available", ha="center", va="center")
        plt.title(title)
        plt.savefig(temp.name)
        plt.close()
        return temp.name

    series = series.fillna(0)

    # =========================
    # LINE CHART (SMART DATES)
    # =========================
    if chart_type == "line":
        ax = series.plot(kind="line")

        if isinstance(series.index, pd.DatetimeIndex):
            locator = mdates.AutoDateLocator(minticks=4, maxticks=8)
            formatter = mdates.ConciseDateFormatter(locator)

            ax.xaxis.set_major_locator(locator)
            ax.xaxis.set_major_formatter(formatter)
            plt.xticks(rotation=25)

    # =========================
    # BAR CHART
    # =========================
    elif chart_type == "bar":
        series.plot(kind="bar")
        plt.xticks(rotation=35, ha="right")

    # =========================
    # PIE CHART (SAFE)
    # =========================
    elif chart_type == "pie":
        series = series[series > 0]

        if len(series) == 0:
            plt.text(0.5, 0.5, "No Data Available", ha="center", va="center")
        else:
            series.plot(kind="pie", autopct="%1.1f%%")

    plt.title(title, fontsize=12, fontweight="bold")
    plt.tight_layout()
    plt.savefig(temp.name)
    plt.close()

    return temp.name


# ==========================================================
# MAIN PDF GENERATOR
# ==========================================================

def generate_pdf(df, title, date_range_str, metrics=None):

    pdf = BusinessPDFReport()
    chart_files = []

    df = df.copy()

    # =========================
    # SAFE COLUMN HANDLING
    # =========================
    if "Date" not in df.columns:
        df["Date"] = pd.to_datetime(datetime.now())

    if "Revenue" not in df.columns:
        df["Revenue"] = df.get("Qty", 0) * 1.0

    if "QR Revenue" not in df.columns:
        df["QR Revenue"] = 0

    if "Drink Profile" not in df.columns:
        df["Drink Profile"] = "Unknown"

    if "Drink Type" not in df.columns:
        df["Drink Type"] = "Unknown"

    if "Qty" not in df.columns:
        df["Qty"] = 0

    # =========================
    # KPI CALCULATIONS
    # =========================
    revenue = float(df["Revenue"].sum())
    qr_revenue = float(df["QR Revenue"].sum())
    cash_revenue = revenue - qr_revenue
    cups = int(df["Qty"].sum())
    avg = revenue / cups if cups > 0 else 0

    try:
        best_seller = (
            df.groupby("Drink Profile")["Qty"]
            .sum()
            .idxmax()
        )
    except Exception:
        best_seller = "N/A"

    # =========================
    # PAGE 1 - EXECUTIVE SUMMARY
    # =========================

    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, title, ln=True)

    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(90, 90, 90)

    pdf.cell(0, 6, f"Reporting Period: {date_range_str}", ln=True)
    pdf.cell(0, 6, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", ln=True)

    pdf.ln(10)

    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 10, "Executive Summary", ln=True)

    pdf.ln(3)

    draw_kpi(pdf, "Total Revenue", f"RM {revenue:,.2f}")
    draw_kpi(pdf, "Cash Revenue", f"RM {cash_revenue:,.2f}")
    draw_kpi(pdf, "QR Revenue", f"RM {qr_revenue:,.2f}")
    draw_kpi(pdf, "Total Cups Sold", cups)
    draw_kpi(pdf, "Avg Per Cup", f"RM {avg:.2f}")
    draw_kpi(pdf, "Best Seller", best_seller)

    # =========================
    # PAGE 2 - DAILY TREND
    # =========================

    if "Revenue" in df.columns:

        daily = df.groupby(df["Date"].dt.date)["Revenue"].sum()

        chart = create_chart(daily, "Daily Revenue Trend", "line")
        chart_files.append(chart)

        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Revenue Trend", ln=True)
        pdf.image(chart, x=10, w=180)

    # =========================
    # PAGE 3 - TOP DRINKS
    # =========================

    top = (
        df.groupby("Drink Profile")["Qty"]
        .sum()
        .sort_values(ascending=False)
        .head(10)
    )

    chart = create_chart(top, "Top Selling Drinks", "bar")
    chart_files.append(chart)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Top Selling Drinks", ln=True)
    pdf.image(chart, x=10, w=180)

    # =========================
    # PAGE 4 - PAYMENT BREAKDOWN
    # =========================

    payment = pd.Series({
        "Cash": cash_revenue,
        "QR": qr_revenue
    })

    chart = create_chart(payment, "Payment Breakdown", "pie")
    chart_files.append(chart)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Payment Breakdown", ln=True)
    pdf.image(chart, x=15, w=160)

    # =========================
    # PAGE 5 - HOT VS COLD
    # =========================

    hot_cold = df.groupby("Drink Type")["Qty"].sum()

    chart = create_chart(hot_cold, "Hot vs Cold Drinks", "pie")
    chart_files.append(chart)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Hot vs Cold Drinks", ln=True)
    pdf.image(chart, x=15, w=160)

    # =========================
    # CLEANUP
    # =========================

    result = bytes(pdf.output())

    for f in chart_files:
        try:
            os.remove(f)
        except Exception:
            pass

    return result
